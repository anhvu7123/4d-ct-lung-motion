#!/usr/bin/env python3
"""Read-only manuscript/data coherence check for the corrected V9 paper.

Usage from the lung_4dct_project directory:
    python verify_manuscript_v9.py
Requires python-docx from requirements.txt. Does not run numerical simulation.
"""
from pathlib import Path
import json
from docx import Document

root=Path(__file__).resolve().parent
j=json.loads((root/'simulation_results.json').read_text(encoding='utf-8'))
doc=Document(root/'manuscript.docx')
sub=j['subvoxel_sensitivity']['targets']
rows={r.cells[0].text.strip():r for r in doc.tables[-1].rows[1:]}
mapping={
 'Isotropic 5 mm':'Isotropic 5mm',
 'ITV + 3 mm':'ITV + 3mm setup',
 'ITV + 4 mm':'ITV + 4mm setup',
 'ITV + 5 mm':'ITV + 5mm setup',
 'Aniso + 2 mm':'Anisotropic (Mid-pos + 2mm setup)',
}
def signed(x):
    if abs(x)<0.0499:x=0
    return f'{x:+.1f}' if x>0 else f'{x:.1f}'
for shown, key in mapping.items():
    assert shown in rows, f'Missing Table 5 row {shown}'
    r=rows[shown]
    dat=sub[key]
    delta=float(r.cells[3].text)
    assert abs(delta-dat['pass_rate_difference_pct'])<0.051,(shown,delta,dat['pass_rate_difference_pct'])
    ci=dat['pass_difference_bootstrap_ci_95_pp']
    assert r.cells[6].text.strip()==f'[{signed(ci[0])}, {signed(ci[1])}]',f'CI mismatch for {shown}'
text='\n'.join(p.text for p in doc.paragraphs)
assert 'raw_fraction_cr_by_patient' in text,'Data availability paragraph is out of date'
assert 'does not include the full 200' not in text,'Prior contradictory data statement still present'
assert 'utilizing CT-based technologies' in text,'AAPM TG-179 title not updated'
assert all('Revised manuscript' not in p.text for s in doc.sections for p in s.footer.paragraphs)
print('PASS: V9 manuscript text, Table 5 paired-bootstrap CIs and JSON agree.')
print('Note: verifies saved outputs and document consistency; does not independently rerun simulations.')
